import csv
import json
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from doc2sheet.config import SUGGESTED_MODELS, Settings
from doc2sheet.export import export_all
from doc2sheet.schema import ExtractedDocument, ExtractionResult, Issue, LineItem
from doc2sheet.service import UserFacingError, resolve_connection, run_batch, summarise
from tests.helpers import FakeLLM, invoice_json


def sample_results():
    doc = ExtractedDocument(
        document_type="invoice",
        vendor_name="Café Zoë",
        document_number="A-1",
        currency="EUR",
        total=Decimal("12.50"),
        line_items=[LineItem(description="Crème brûlée", quantity=Decimal("1"), amount=Decimal("12.50"))],
    )
    return [
        ExtractionResult(source="ok.pdf", document=doc),
        ExtractionResult(
            source="review.pdf",
            document=doc,
            issues=[Issue(level="error", code="total_mismatch", message="Totals don't add up")],
        ),
        ExtractionResult(source="bad.pdf", error="Model call failed: 401"),
    ]


def test_export_all_writes_three_formats(tmp_path):
    files = export_all(sample_results(), tmp_path)

    wb = load_workbook(files["xlsx"])
    assert wb.sheetnames == ["Documents", "Line items", "Issues"]
    docs = list(wb["Documents"].values)
    assert docs[1][:2] == ("ok.pdf", "ok")
    assert docs[1][docs[0].index("total")] == 12.5
    assert docs[3][1] == "failed"
    assert len(list(wb["Line items"].values)) == 3  # header + one item for each successful doc
    assert [row[2] for row in list(wb["Issues"].values)[1:]] == ["total_mismatch", "failed"]

    raw = files["csv"].read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    rows = list(csv.DictReader(files["csv"].open(encoding="utf-8-sig")))
    assert rows[0]["vendor"] == "Café Zoë" and rows[2]["issues"].startswith("Model call failed")

    data = json.loads(files["json"].read_text(encoding="utf-8"))
    assert data[0]["document"]["total"] == 12.5 and data[1]["status"] == "review"
    assert "raw_response" not in data[0]


def test_summary_counts():
    text = summarise(sample_results())
    assert "3 document(s)" in text and "1 passed" in text and "1 need review" in text and "1 failed" in text


def test_run_batch(samples, tmp_path):
    out = run_batch([samples / "invoice_pixel_pine_eur.pdf"], FakeLLM([invoice_json()]), out_dir=tmp_path)
    assert out.results[0].status in {"ok", "review"}
    assert out.documents[0][0] == "invoice_pixel_pine_eur.pdf"
    assert len(out.line_items) == 4
    assert set(out.files) == {"xlsx", "csv", "json"}


# -- credential rules for the public demo -------------------------------------

SERVER = Settings(base_url="https://router.huggingface.co/v1", api_key="server-token", model=SUGGESTED_MODELS[0])


def test_server_key_is_used_for_the_server_endpoint():
    conn = resolve_connection(SERVER, model="", api_key="", base_url="", public_demo=True)
    assert conn.api_key == "server-token" and conn.model == SUGGESTED_MODELS[0]


def test_server_key_is_never_sent_to_another_endpoint():
    conn = resolve_connection(SERVER, model="m", api_key="", base_url="https://evil.example/v1", public_demo=False)
    assert conn.api_key is None and conn.base_url == "https://evil.example/v1"


def test_public_demo_cannot_change_endpoint():
    with pytest.raises(UserFacingError):
        resolve_connection(SERVER, model="", api_key="mine", base_url="https://evil.example/v1", public_demo=True)


def test_public_demo_limits_models_unless_user_brings_a_key():
    with pytest.raises(UserFacingError, match="Custom models"):
        resolve_connection(SERVER, model="some/huge-model", api_key="", base_url="", public_demo=True)
    conn = resolve_connection(SERVER, model="some/huge-model", api_key="mine", base_url="", public_demo=True)
    assert conn.api_key == "mine"


def test_missing_hf_token_is_explained():
    no_key = Settings(base_url="https://router.huggingface.co/v1", api_key=None)
    with pytest.raises(UserFacingError, match="Hugging Face token"):
        resolve_connection(no_key, model="", api_key="", base_url="", public_demo=False)


def test_local_ollama_needs_no_key():
    conn = resolve_connection(
        SERVER, model="gemma3:4b", api_key="", base_url="http://localhost:11434/v1/", public_demo=False
    )
    assert conn.base_url == "http://localhost:11434/v1" and conn.api_key is None
