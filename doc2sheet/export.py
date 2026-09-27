"""Flatten results into rows and write them to Excel, CSV and JSON."""

from __future__ import annotations

import csv
import json
from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from doc2sheet.schema import ExtractionResult

DOCUMENT_COLUMNS = [
    "file",
    "status",
    "type",
    "vendor",
    "vendor_tax_id",
    "customer",
    "number",
    "issue_date",
    "due_date",
    "currency",
    "subtotal",
    "discount",
    "other_charges",
    "tax",
    "total",
    "line_items",
    "issues",
]
LINE_ITEM_COLUMNS = [
    "file",
    "number",
    "vendor",
    "currency",
    "line",
    "description",
    "quantity",
    "unit_price",
    "amount",
]
ISSUE_COLUMNS = ["file", "level", "code", "message"]
MONEY_COLUMNS = {"subtotal", "discount", "other_charges", "tax", "total", "unit_price", "amount"}


def document_rows(results: Sequence[ExtractionResult]) -> list[dict[str, Any]]:
    rows = []
    for r in results:
        d = r.document
        row: dict[str, Any] = {c: None for c in DOCUMENT_COLUMNS}
        row.update(file=r.source, status=r.status)
        if d is not None:
            row.update(
                type=d.document_type,
                vendor=d.vendor_name,
                vendor_tax_id=d.vendor_tax_id,
                customer=d.customer_name,
                number=d.document_number,
                issue_date=d.issue_date,
                due_date=d.due_date,
                currency=d.currency,
                subtotal=d.subtotal,
                discount=d.discount,
                other_charges=d.other_charges,
                tax=d.tax,
                total=d.total,
                line_items=len(d.line_items),
            )
        messages = [i.message for i in r.issues]
        if r.error:
            messages.insert(0, r.error)
        row["issues"] = " | ".join(messages) or None
        rows.append(row)
    return rows


def line_item_rows(results: Sequence[ExtractionResult]) -> list[dict[str, Any]]:
    rows = []
    for r in results:
        if r.document is None:
            continue
        d = r.document
        for n, item in enumerate(d.line_items, start=1):
            rows.append(
                dict(
                    file=r.source,
                    number=d.document_number,
                    vendor=d.vendor_name,
                    currency=d.currency,
                    line=n,
                    description=item.description,
                    quantity=item.quantity,
                    unit_price=item.unit_price,
                    amount=item.amount,
                )
            )
    return rows


def issue_rows(results: Sequence[ExtractionResult]) -> list[dict[str, Any]]:
    rows = []
    for r in results:
        if r.error:
            rows.append(dict(file=r.source, level="error", code="failed", message=r.error))
        for issue in r.issues:
            rows.append(dict(file=r.source, level=issue.level, code=issue.code, message=issue.message))
    return rows


def as_table(rows: list[dict[str, Any]], columns: list[str]) -> list[list[Any]]:
    """Rows as plain lists with display-friendly values (for UI tables)."""
    return [[_display(row.get(c)) for c in columns] for row in rows]


def _display(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else value


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def write_json(results: Sequence[ExtractionResult], path: str | Path) -> Path:
    path = Path(path)
    data = [r.model_dump(mode="json", exclude={"raw_response"}) for r in results]
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def write_csv(results: Sequence[ExtractionResult], path: str | Path) -> Path:
    """One row per document. UTF-8 with BOM so Excel on Windows opens accents correctly."""
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=DOCUMENT_COLUMNS)
        writer.writeheader()
        for row in document_rows(results):
            writer.writerow({k: _csv_value(v) for k, v in row.items()})
    return path


def _csv_value(value: Any) -> Any:
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        # 630 -> "630.00", but keep extra precision such as 0.125
        return f"{value:.2f}" if value == value.quantize(Decimal("0.01")) else str(value)
    return "" if value is None else value


def write_xlsx(results: Sequence[ExtractionResult], path: str | Path) -> Path:
    """A reviewer-friendly workbook: Documents, Line items and Issues sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    status_fills = {
        "ok": PatternFill("solid", fgColor="D8F3DC"),
        "review": PatternFill("solid", fgColor="FFF3BF"),
        "failed": PatternFill("solid", fgColor="FFD6D6"),
        "error": PatternFill("solid", fgColor="FFD6D6"),
        "warning": PatternFill("solid", fgColor="FFF3BF"),
    }

    wb = Workbook()
    sheets = [
        ("Documents", DOCUMENT_COLUMNS, document_rows(results), "status"),
        ("Line items", LINE_ITEM_COLUMNS, line_item_rows(results), None),
        ("Issues", ISSUE_COLUMNS, issue_rows(results), "level"),
    ]
    for index, (title, columns, rows, colour_by) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in rows:
            ws.append([_xlsx_value(row.get(c)) for c in columns])

        widths = {c: len(c) for c in columns}
        for col_idx, column in enumerate(columns, start=1):
            letter = get_column_letter(col_idx)
            for cell in ws[letter][1:]:
                if column in MONEY_COLUMNS:
                    cell.number_format = "#,##0.00"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(wrap_text=column in ("issues", "message", "description"), vertical="top")
                if colour_by == column and cell.value in status_fills:
                    cell.fill = status_fills[cell.value]
                widths[column] = max(widths[column], len(str(cell.value or "")))
            minimum = 12 if column in MONEY_COLUMNS else 0
            ws.column_dimensions[letter].width = min(max(widths[column] + 2, minimum), 60)

        ws.freeze_panes = "A2"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if rows:
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return Path(path)


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_all(results: Sequence[ExtractionResult], out_dir: str | Path, stem: str = "doc2sheet") -> dict[str, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    return {
        "xlsx": write_xlsx(results, out_dir / f"{stem}.xlsx"),
        "csv": write_csv(results, out_dir / f"{stem}.csv"),
        "json": write_json(results, out_dir / f"{stem}.json"),
    }
